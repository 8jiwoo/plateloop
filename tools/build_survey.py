"""
PlateLoop survey kit: questionnaires for four groups, DEMO responses, an Excel analysis workbook and a findings report.

  pip install openpyxl pandas selenium
  python tools/build_survey.py

Writes to docs/survey/:
  plateloop-questionnaires.pdf            printable questionnaires (real, ready to use)
  plateloop-survey-analysis-DEMO.xlsx     raw demo responses, formula-driven analysis, charts and key findings
  plateloop-key-findings-DEMO.pdf         the findings as a short report

ALL RESPONSES ARE SYNTHETIC (randomly generated with a fixed seed) to demonstrate the analysis.
They are not real survey results. Replace them with real responses before drawing conclusions.
"""
import base64, os, random, sys, time
from decimal import Decimal, ROUND_HALF_UP
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'docs', 'survey')
SEED = 2026
MAXROW = 1000  # analysis formulas cover rows 2..1000 so real data can be pasted in

L5 = {'agree': ['1 Strongly disagree', '2', '3', '4', '5 Strongly agree']}
# ================================================================ questionnaires (single source of truth)
# question: (code, text, type, options or None, demo distribution)
#   type 'cat'   -> one option; dist = weights per option
#   type 'lik'   -> 1-5 scale; dist = weights for 1..5; options = (low label, high label)
#   type 'num'   -> number; dist = (mean, sd, min, max)
#   type 'text'  -> free text; dist = list of demo comments
GROUPS = [
 dict(key='Students', sheet='Students_Data', n=120, title='Students (secondary school and university)',
      intro='About your school or campus canteen lunch. Takes about 4 minutes. Anonymous.',
      qs=[
   ('S1', 'Which are you?', 'cat', ['Secondary school', 'University'], [.58, .42]),
   ('S2', 'How often do you leave food on your tray?', 'cat', ['Never', 'Rarely', 'Sometimes', 'Often', 'Almost always'], [.08, .20, .37, .25, .10]),
   ('S3', 'What is the main reason you leave food?', 'cat', ['Portion too big', 'Don’t like the taste', 'Not hungry', 'Ran out of time', 'Didn’t recognise the dish'], [.38, .27, .15, .12, .08]),
   ('S4', 'Would you take a smaller portion if you could ask for one easily?', 'cat', ['Yes', 'Maybe', 'No'], [.58, .27, .15]),
   ('S5', 'Are you told to finish everything on your plate?', 'cat', ['Yes', 'Sometimes', 'No'], None),
   ('S6', 'I would like to see what I ate at lunch (food groups and nutrients).', 'lik', ('Strongly disagree', 'Strongly agree'), [.06, .10, .24, .34, .26]),
   ('S7', 'Which would you rather see about your lunch?', 'cat', ['Food groups', 'Calories', 'Both', 'Neither'], [.46, .18, .22, .14]),
   ('S8', 'I would scan my tray before and after lunch if it fed a virtual pet.', 'lik', ('Strongly disagree', 'Strongly agree'), None),
   ('S9', 'A weekly class competition on food waste would motivate me.', 'lik', ('Strongly disagree', 'Strongly agree'), [.07, .10, .22, .33, .28]),
   ('S10', 'I would be comfortable signing in to a scanner with my face.', 'lik', ('Strongly disagree', 'Strongly agree'), [.14, .17, .25, .25, .19]),
   ('S11', 'Anything else about lunch or food waste at your canteen?', 'text', None, [
      'The rice portion is always too big for me.', 'I would eat more vegetables if they tasted better.', 'Sometimes I don’t know what the dish is.',
      'Queues are too long so I rush my food.', 'A pet game would be fun if it is not too childish.', 'I don’t want my name on a leaderboard.',
      'More fruit options please.', 'I prefer not to use face scanning.', 'Being told to finish everything makes me feel bad.', ''])]),
 dict(key='Kitchen', sheet='Kitchen_Data', n=24, title='Canteen and kitchen staff',
      intro='About how your canteen plans food and handles leftovers. Takes about 4 minutes.',
      qs=[
   ('K1', 'What is your role?', 'cat', ['Canteen operator', 'Cook', 'Kitchen assistant'], [.25, .42, .33]),
   ('K2', 'How do you mainly decide how much to cook each day?', 'cat', ['Past experience', 'Headcount × standard portion', 'Previous day’s leftovers', 'Written records'], [.50, .33, .12, .05]),
   ('K3', 'Roughly what percentage of the food you cook is thrown away on a typical day?', 'num', None, (22, 7, 5, 45)),
   ('K4', 'Do you know which dishes come back on trays most?', 'cat', ['Yes, accurately', 'Roughly', 'No'], [.17, .58, .25]),
   ('K5', 'How often do you prepare more food than is eaten?', 'cat', ['Never', 'Rarely', 'Weekly', 'Several times a week', 'Daily'], [.04, .13, .38, .33, .12]),
   ('K6', 'A suggested order for tomorrow, based on what was eaten, would be useful.', 'lik', ('Not useful', 'Very useful'), [.02, .05, .12, .36, .45]),
   ('K7', 'Seeing waste broken down by dish would be useful.', 'lik', ('Not useful', 'Very useful'), [.02, .04, .10, .35, .49]),
   ('K8', 'A monthly carbon report would be useful.', 'lik', ('Not useful', 'Very useful'), [.10, .18, .32, .25, .15]),
   ('K9', 'How many minutes a day do you spend planning orders and quantities?', 'num', None, (35, 12, 10, 75)),
   ('K10', 'I am worried a new system would add to my workload.', 'lik', ('Strongly disagree', 'Strongly agree'), [.08, .14, .26, .32, .20]),
   ('K11', 'Anything else?', 'text', None, [
      'We cook the same amount every day regardless of attendance.', 'Vegetables come back the most.', 'Exam weeks and trips change everything.',
      'I don’t have time to fill in forms.', 'If it saves money the boss will like it.', 'Soup is always wasted.', ''])]),
 dict(key='Hospital', sheet='Hospital_Data', n=36, title='Hospital staff (nurses, dietitians, ward managers)',
      intro='About how patients’ food intake is tracked on your ward. Takes about 4 minutes.',
      qs=[
   ('H1', 'What is your role?', 'cat', ['Nurse', 'Dietitian', 'Ward manager'], [.67, .22, .11]),
   ('H2', 'How is patients’ food intake mainly tracked on your ward?', 'cat', ['Food charts', 'Visual estimate', 'Patient report', 'Not tracked'], [.55, .25, .14, .06]),
   ('H3', 'I am confident our intake records are accurate.', 'lik', ('Strongly disagree', 'Strongly agree'), [.17, .36, .30, .14, .03]),
   ('H4', 'How many minutes per shift do you spend documenting food intake?', 'num', None, (28, 10, 5, 60)),
   ('H5', 'How quickly is a patient who stops eating usually noticed?', 'cat', ['Same day', 'After 2–3 days', 'After 4+ days', 'Often missed'], [.14, .44, .25, .17]),
   ('H6', 'What is the most common reason patients leave food?', 'cat', ['Poor appetite', 'Taste', 'Portion too big', 'Too tired', 'Pain or nausea'], [.36, .19, .17, .12, .16]),
   ('H7', 'Automatic intake records for every tray would be useful.', 'lik', ('Not useful', 'Very useful'), [.01, .03, .08, .30, .58]),
   ('H8', 'An alert that includes the patient’s reason for not eating would be useful.', 'lik', ('Not useful', 'Very useful'), [.02, .04, .10, .32, .52]),
   ('H9', 'Would patients benefit from choosing their next day’s meals and portion size?', 'cat', ['Yes', 'Maybe', 'No'], [.64, .28, .08]),
   ('H10', 'Anything else?', 'text', None, [
      'Food charts are filled in at the end of the shift from memory.', 'Elderly patients often leave most of their tray.', 'We need to know protein intake for wound healing.',
      'Alerts must not create more paperwork.', 'Patients would like more choice.', ''])]),
 dict(key='Office', sheet='Office_Data', n=80, title='Office workers',
      intro='About lunch at your workplace canteen. Takes about 4 minutes. Anonymous.',
      qs=[
   ('O1', 'Your age group', 'cat', ['20–29', '30–39', '40–49', '50+'], [.30, .35, .22, .13]),
   ('O2', 'How many days a week do you eat at the office canteen?', 'num', None, None),
   ('O3', 'Do you have a health goal for your meals?', 'cat', ['Lose weight', 'Build muscle', 'Steady energy', 'Eat balanced', 'None'], [.30, .14, .20, .21, .15]),
   ('O4', 'I know the nutrients in my canteen lunch.', 'lik', ('Strongly disagree', 'Strongly agree'), [.28, .38, .20, .10, .04]),
   ('O5', 'How often do you feel an energy dip in the afternoon?', 'cat', ['Never', 'Rarely', 'Sometimes', 'Often', 'Daily'], [.05, .15, .35, .32, .13]),
   ('O6', 'Would you pre-order lunch to skip the queue?', 'cat', ['Yes', 'Maybe', 'No'], [.56, .29, .15]),
   ('O7', 'A daily lunch suggestion for my health goal would be useful.', 'lik', ('Not useful', 'Very useful'), [.04, .08, .20, .38, .30]),
   ('O8', 'Would you share your meal data with your employer?', 'cat', ['Yes', 'Only anonymous totals', 'No'], [.16, .58, .26]),
   ('O9', 'How many minutes do you usually queue at lunch?', 'num', None, (11, 4, 2, 25)),
   ('O10', 'Anything else?', 'text', None, [
      'Healthy options run out early.', 'I never know how many calories are in the mixed rice.', 'Pre-ordering would save me 15 minutes.',
      'I don’t want HR to see what I eat.', 'Afternoon slump after noodles.', ''])]),
]

# ================================================================ demo data
def make_data():
    rnd = random.Random(SEED); data = {}
    for G in GROUPS:
        rows = []
        for i in range(G['n']):
            r = {'ID': f"{G['key'][0]}{i+1:03d}"}
            for code, text, typ, opts, dist in G['qs']:
                if typ == 'cat':
                    if code == 'S5':  # being told to finish depends on level
                        w = [.45, .35, .20] if r['S1'] == 'Secondary school' else [.10, .25, .65]
                    else: w = dist
                    r[code] = rnd.choices(opts, w)[0]
                elif typ == 'lik':
                    if code == 'S8':  # the pet appeals more to secondary students
                        w = [.03, .06, .16, .35, .40] if r['S1'] == 'Secondary school' else [.10, .16, .30, .28, .16]
                    else: w = dist
                    r[code] = rnd.choices([1, 2, 3, 4, 5], w)[0]
                elif typ == 'num':
                    if code == 'O2': r[code] = rnd.choices([0, 1, 2, 3, 4, 5], [.02, .05, .13, .25, .30, .25])[0]
                    else:
                        m, sd, lo, hi = dist; r[code] = int(min(hi, max(lo, round(rnd.gauss(m, sd)))))
                else:
                    r[code] = rnd.choice(dist) if rnd.random() < .45 else ''
            rows.append(r)
        data[G['key']] = pd.DataFrame(rows)
    return data


def tidy(ch, pct=True):
    """Clean chart: one green colour, value-only labels, visible axes, answers top to bottom."""
    from openpyxl.chart.shapes import GraphicalProperties
    ch.varyColors = False
    s = ch.series[0]; s.graphicalProperties = GraphicalProperties(solidFill='1E7A44'); s.graphicalProperties.line.solidFill = '1E7A44'
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
    ch.dataLabels.showSerName = False; ch.dataLabels.showCatName = False; ch.dataLabels.showLegendKey = False; ch.dataLabels.showPercent = False
    ch.x_axis.delete = False; ch.y_axis.delete = True  # values are labelled on the bars
    ch.x_axis.scaling.orientation = 'maxMin'
    ch.y_axis.majorGridlines = None
    if pct: ch.y_axis.numFmt = '0%'; ch.y_axis.scaling.min = 0
    ch.gapWidth = 60
    return ch

# ================================================================ workbook
FONT = 'Arial'
F = lambda **k: Font(name=FONT, **k)
HEAD = PatternFill('solid', fgColor='1E7A44'); SUB = PatternFill('solid', fgColor='E3F2E8'); WARN = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='D0D5D1'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

def col_of(G, code):
    cols = ['ID'] + [q[0] for q in G['qs']]
    return get_column_letter(cols.index(code) + 1)

def rng(G, code):
    c = col_of(G, code); return f"'{G['sheet']}'!${c}$2:${c}${MAXROW}"

def build_workbook(data, path):
    wb = Workbook()
    # ---- README
    ws = wb.active; ws.title = 'README'
    ws['A1'] = 'PlateLoop survey analysis'; ws['A1'].font = F(bold=True, size=16)
    ws['A2'] = 'DEMO DATA: every response in this workbook is synthetic (randomly generated) to show how the analysis works. These are NOT real survey results.'
    ws['A2'].font = F(bold=True, color='C00000'); ws['A2'].fill = WARN
    lines = [('', ''), ('Sheet', 'What it contains'),
             ('Questionnaires', 'Every question, its code, type and answer options, by group.'),
             ('Students_Data … Office_Data', 'One row per respondent (demo). Blue text = input data. Paste real responses here, same columns.'),
             ('Analysis', 'Counts, shares, averages and segment comparisons. All formulas: they update when the data changes (rows 2–1000).'),
             ('Key_Findings', 'Findings written as formulas from the Analysis sheet, each linked to a PlateLoop feature.'),
             ('', ''), ('Legend', ''), ('Blue text', 'Input data (responses)'), ('Black text', 'Formulas'),
             ('Likert questions', '1 = Strongly disagree / Not useful … 5 = Strongly agree / Very useful. "Agree" = share answering 4 or 5.'),
             ('Demo sample sizes', ', '.join(f"{G['key']}: {G['n']}" for G in GROUPS) + f' (seed {SEED}).')]
    for i, (a, b) in enumerate(lines, start=3):
        ws.cell(i, 1, a).font = F(bold=a in ('Sheet', 'Legend')); ws.cell(i, 2, b).font = F()
    ws['A9'].font = F(bold=True); ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 110
    ws['A11'].font = F(color='0000FF')
    # ---- Questionnaires
    q = wb.create_sheet('Questionnaires')
    for j, h in enumerate(['Group', 'Code', 'Question', 'Type', 'Options / scale'], 1):
        c = q.cell(1, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = HEAD
    r = 2
    for G in GROUPS:
        for code, text, typ, opts, _ in G['qs']:
            o = '; '.join(opts) if typ == 'cat' else (f'1 = {opts[0]} … 5 = {opts[1]}' if typ == 'lik' else 'Number' if typ == 'num' else 'Free text')
            for j, v in enumerate([G['title'], code, text, {'cat': 'Single choice', 'lik': 'Scale 1–5', 'num': 'Number', 'text': 'Open'}[typ], o], 1):
                c = q.cell(r, j, v); c.font = F(); c.alignment = Alignment(wrap_text=True, vertical='top')
            r += 1
    for col, w in zip('ABCDE', (30, 7, 60, 14, 60)): q.column_dimensions[col].width = w
    q.freeze_panes = 'A2'
    # ---- data sheets
    for G in GROUPS:
        d = wb.create_sheet(G['sheet']); df = data[G['key']]
        for j, h in enumerate(df.columns, 1):
            c = d.cell(1, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = HEAD
            qtext = next((x[1] for x in G['qs'] if x[0] == h), 'Respondent ID')
            from openpyxl.comments import Comment
            c.comment = Comment(qtext, 'PlateLoop')
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                c = d.cell(i, j, v); c.font = F(color='0000FF')
        for j in range(1, len(df.columns) + 1): d.column_dimensions[get_column_letter(j)].width = 24 if j > 1 else 8
        d.freeze_panes = 'B2'
    # ---- Analysis
    a = wb.create_sheet('Analysis'); A = {}
    a['A1'] = 'Analysis (DEMO DATA)'; a['A1'].font = F(bold=True, size=14)
    a['A2'] = 'All values are formulas over the data sheets (rows 2–1000).'; a['A2'].font = F(italic=True, color='666666')
    r = 4; charts = []
    for G in GROUPS:
        a.cell(r, 1, G['title']).font = F(bold=True, size=13, color='1E7A44')
        a.cell(r + 1, 1, 'Respondents'); a.cell(r + 1, 2, f"=COUNTA('{G['sheet']}'!$A$2:$A${MAXROW})")
        A[(G['key'], 'n')] = f'Analysis!$B${r+1}'
        for c in (a.cell(r + 1, 1), a.cell(r + 1, 2)): c.font = F(bold=True)
        r += 3
        lik_rows = []
        for code, text, typ, opts, _ in G['qs']:
            R = rng(G, code)
            if typ == 'cat':
                a.cell(r, 1, f'{code}  {text}').font = F(bold=True)
                for j, h in enumerate(['Answer', 'Count', 'Share'], 1):
                    c = a.cell(r + 1, j, h); c.font = F(bold=True); c.fill = SUB
                top = r + 2
                for k, o in enumerate(opts):
                    rr = top + k
                    a.cell(rr, 1, o).font = F()
                    a.cell(rr, 2, f'=COUNTIF({R},A{rr})').font = F()
                    c = a.cell(rr, 3, f'=IFERROR(B{rr}/SUM($B${top}:$B${top+len(opts)-1}),0)'); c.font = F(); c.number_format = '0%'
                    A[(code, o)] = f'Analysis!$C${rr}'
                    A[(code, o, 'n')] = f'Analysis!$B${rr}'
                ch = BarChart(); ch.type = 'bar'; ch.style = 10; ch.title = f'{code} {text}'[:60]; ch.legend = None
                ch.add_data(Reference(a, min_col=3, min_row=top, max_row=top + len(opts) - 1), titles_from_data=False)
                ch.set_categories(Reference(a, min_col=1, min_row=top, max_row=top + len(opts) - 1))
                tidy(ch)
                ch.height = 5.5 + .45 * len(opts); ch.width = 14
                charts.append((ch, f'F{r}'))
                r = top + len(opts) + 1
                r = max(r, r)  # spacing
                r += max(0, int(ch.height * 2) - len(opts) - 2)
            elif typ == 'lik':
                lik_rows.append((code, text, R))
            elif typ == 'num':
                a.cell(r, 1, f'{code}  {text}').font = F(bold=True)
                for k, (lab, f) in enumerate([('Average', f'=IFERROR(AVERAGE({R}),0)'), ('Median', f'=IFERROR(MEDIAN({R}),0)'), ('Minimum', f'=IFERROR(MIN({R}),0)'), ('Maximum', f'=IFERROR(MAX({R}),0)')]):
                    a.cell(r + 1 + k, 1, lab).font = F(); c = a.cell(r + 1 + k, 2, f); c.font = F(); c.number_format = '0.0'
                    A[(code, lab)] = f'Analysis!$B${r+1+k}'
                r += 7
        if lik_rows:
            a.cell(r, 1, 'Scale questions (1–5)').font = F(bold=True)
            for j, h in enumerate(['Question', 'Average (1–5)', 'Agree (4–5)', 'Answers'], 1):
                c = a.cell(r + 1, j, h); c.font = F(bold=True); c.fill = SUB
            top = r + 2
            for k, (code, text, R) in enumerate(lik_rows):
                rr = top + k
                a.cell(rr, 1, f'{code}  {text}').font = F()
                c = a.cell(rr, 2, f'=IFERROR(AVERAGE({R}),0)'); c.font = F(); c.number_format = '0.0'
                c = a.cell(rr, 3, f'=IFERROR(COUNTIF({R},">=4")/COUNT({R}),0)'); c.font = F(); c.number_format = '0%'
                a.cell(rr, 4, f'=COUNT({R})').font = F()
                A[(code, 'avg')] = f'Analysis!$B${rr}'; A[(code, 'agree')] = f'Analysis!$C${rr}'
            ch = BarChart(); ch.type = 'bar'; ch.style = 10; ch.title = f"{G['key']}: share who agree (4–5)"; ch.legend = None
            ch.add_data(Reference(a, min_col=3, min_row=top, max_row=top + len(lik_rows) - 1), titles_from_data=False)
            ch.set_categories(Reference(a, min_col=1, min_row=top, max_row=top + len(lik_rows) - 1))
            tidy(ch); ch.y_axis.scaling.max = 1
            ch.height = 7; ch.width = 14; charts.append((ch, f'F{r}'))
            r = top + len(lik_rows) + 10
        r += 2
    # segment comparisons
    a.cell(r, 1, 'Segment comparisons').font = F(bold=True, size=13, color='1E7A44'); r += 1
    S = next(G for G in GROUPS if G['key'] == 'Students')
    segs = [('Would scan to feed a virtual pet (S8), average: secondary', f"=IFERROR(AVERAGEIFS({rng(S,'S8')},{rng(S,'S1')},\"Secondary school\"),0)", 'S8sec', '0.0'),
            ('Would scan to feed a virtual pet (S8), average: university', f"=IFERROR(AVERAGEIFS({rng(S,'S8')},{rng(S,'S1')},\"University\"),0)", 'S8uni', '0.0'),
            ('Told to finish their plate (S5 = Yes): secondary', f"=IFERROR(COUNTIFS({rng(S,'S5')},\"Yes\",{rng(S,'S1')},\"Secondary school\")/COUNTIF({rng(S,'S1')},\"Secondary school\"),0)", 'S5sec', '0%'),
            ('Told to finish their plate (S5 = Yes): university', f"=IFERROR(COUNTIFS({rng(S,'S5')},\"Yes\",{rng(S,'S1')},\"University\")/COUNTIF({rng(S,'S1')},\"University\"),0)", 'S5uni', '0%')]
    for lab, f, key, fmt in segs:
        a.cell(r, 1, lab).font = F(); c = a.cell(r, 2, f); c.font = F(); c.number_format = fmt
        A[key] = f'Analysis!$B${r}'; r += 1
    for ch, anchor in charts: a.add_chart(ch, anchor)
    a.column_dimensions['A'].width = 62; a.column_dimensions['B'].width = 14; a.column_dimensions['C'].width = 12; a.column_dimensions['D'].width = 10
    # ---- Key findings (formulas)
    k = wb.create_sheet('Key_Findings')
    k['A1'] = 'Key findings (DEMO DATA: synthetic responses, not real results)'; k['A1'].font = F(bold=True, size=14, color='C00000')
    for j, h in enumerate(['#', 'Group', 'Finding (calculated from the Analysis sheet)', 'What it means for PlateLoop'], 1):
        c = k.cell(3, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = HEAD
    P = lambda key, fmt='0%': f'TEXT({A[key]},"{fmt}")'
    findings = [
     ('Students', f'="Portion size is the top reason for leaving food: "&{P(("S3","Portion too big"))}&" chose it, and "&{P(("S4","Yes"))}&" would take a smaller portion if they could ask easily."', 'The before scan and right-portion rewards in Loopi: a small plate finished counts fully.'),
     ('Students', f'="More students prefer food groups ("&{P(("S7","Food groups"))}&") than calories ("&{P(("S7","Calories"))}&") when seeing what they ate."', 'Loopi shows food groups by default; calories stay optional.'),
     ('Students', f'="A virtual pet appeals more to secondary students (average "&{P("S8sec","0.0")}&"/5) than to university students ("&{P("S8uni","0.0")}&"/5)."', 'Lead with the Loopi pet in schools; emphasise nutrition feedback on campuses.'),
     ('Students', f'="Only "&{P(("S10","agree"))}&" are comfortable with face sign-in."', 'Keep the class-and-number alternative and make face sign-in optional.'),
     ('Kitchen', f'="Most staff plan by experience or headcount ("&TEXT({A[("K2","Past experience")]}+{A[("K2","Headcount × standard portion")]},"0%")&"), and only "&{P(("K4","Yes, accurately"))}&" know accurately which dishes come back."', 'Kitchen dashboard: waste by dish and a forecast from real eating data.'),
     ('Kitchen', f'="Staff estimate "&{P(("K3","Average"),"0")}&"% of cooked food is thrown away, and spend "&{P(("K9","Average"),"0")}&" minutes a day planning quantities."', 'Tomorrow’s suggested order saves planning time and cuts over-preparation.'),
     ('Kitchen', f'="A suggested order is rated useful by "&{P(("K6","agree"))}&" and waste by dish by "&{P(("K7","agree"))}&", but "&{P(("K10","agree"))}&" worry about extra workload."', 'Automatic measurement only: no forms or manual entry for kitchen staff.'),
     ('Hospital', f'="Only "&{P(("H3","agree"))}&" of hospital staff trust their intake records, and staff spend "&{P(("H4","Average"),"0")}&" minutes a shift documenting intake."', 'Loopi Care records every tray automatically.'),
     ('Hospital', f'="Patients who stop eating are noticed the same day in only "&{P(("H5","Same day"))}&" of cases."', 'Nurse-station alerts after poor meals, the same day.'),
     ('Hospital', f'="Automatic intake records are rated useful by "&{P(("H7","agree"))}&"; "&{P(("H9","Yes"))}&" say patients would benefit from choosing tomorrow’s meals."', 'Loopi Care: per-tray intake and patient meal choice.'),
     ('Office', f'="Only "&{P(("O4","agree"))}&" of office workers know the nutrients in their lunch, and "&TEXT({A[("O5","Often")]}+{A[("O5","Daily")]},"0%")&" feel an afternoon energy dip often or daily."', 'Loopi Work: nutrients per lunch and a 3 pm energy check-in.'),
     ('Office', f'="{"{}"}"&{P(("O6","Yes"))}&" would pre-order lunch to skip a typical "&{P(("O9","Average"),"0")}&"-minute queue."', 'Pre-ordering in Loopi Work; kitchens cook to orders.'),
     ('Office', f'="Only "&{P(("O8","Yes"))}&" would share meal data with their employer; "&{P(("O8","Only anonymous totals"))}&" would share anonymous totals only."', 'Private by default: employers see team totals, never names.'),
    ]
    for i, (grp, f, mean) in enumerate(findings, 4):
        k.cell(i, 1, i - 3).font = F(bold=True)
        k.cell(i, 2, grp).font = F(bold=True)
        c = k.cell(i, 3, f.replace('="{}"&', '=')); c.font = F(); c.alignment = Alignment(wrap_text=True, vertical='top')
        c = k.cell(i, 4, mean); c.font = F(color='1E7A44'); c.alignment = Alignment(wrap_text=True, vertical='top')
        k.row_dimensions[i].height = 48
    for col, w in zip('ABCD', (4, 11, 80, 55)): k.column_dimensions[col].width = w
    k.freeze_panes = 'A4'
    wb.move_sheet('Key_Findings', offset=-(len(wb.sheetnames) - 2))
    for s in wb.worksheets:
        s.sheet_view.showGridLines = s.title not in ('README', 'Key_Findings')
    wb.save(path)
    return A, findings

# ================================================================ PDFs (questionnaires + findings)
def P(v):
    """Percent rounded half-up, the same way Excel's TEXT(x,"0%") does, so the PDF and the workbook agree."""
    return f'{int(Decimal(str(v * 100)).quantize(Decimal(0), rounding=ROUND_HALF_UP))}%'

CSS = '''@page{size:A4;margin:16mm 16mm 16mm}*{box-sizing:border-box}body{margin:0;font-family:Inter,Arial,sans-serif;color:#16201A;font-size:10pt;line-height:1.45;-webkit-print-color-adjust:exact;print-color-adjust:exact}
h1{font-size:20pt;margin:0 0 4px}h2{font-size:13pt;margin:0 0 4px;color:#1E7A44}.sub{color:#4B5550;margin:0 0 14px}
.grp{page-break-before:always}.grp:first-of-type{page-break-before:auto}
.q{border-top:1px solid #E3E6E3;padding:9px 0 8px;page-break-inside:avoid}.q b{display:block;margin-bottom:6px}.q b span{color:#1E7A44;margin-right:6px}
.opts{display:flex;flex-wrap:wrap;gap:6px 18px;font-size:9.5pt;color:#4B5550}.opts i{display:inline-block;width:11px;height:11px;border:1.4px solid #4B5550;border-radius:50%;margin-right:5px;vertical-align:-1px}
.scale{display:flex;gap:10px;align-items:center;font-size:9pt;color:#8A928D}.scale i{display:inline-grid;place-items:center;width:22px;height:22px;border:1.4px solid #4B5550;border-radius:50%;font-style:normal;color:#16201A;font-size:9pt}
.line{border-bottom:1px solid #B9C0BB;height:22px;width:60%}.box{border:1px solid #B9C0BB;height:48px}
.demo{background:#FFF2CC;border-left:4px solid #C9741C;padding:8px 12px;margin:0 0 14px;font-size:9.5pt}
.fd{display:grid;grid-template-columns:28px 1fr;gap:10px;border-top:1px solid #E3E6E3;padding:9px 0;page-break-inside:avoid}.fd .n{font-weight:800;color:#1E7A44}.fd em{display:block;font-style:normal;color:#1E7A44;font-size:9pt;margin-top:3px}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:12px 22px;margin:10px 0}.cap{font-size:8.5pt;color:#8A928D;margin:2px 0 8px}
.kf{display:grid;grid-template-columns:repeat(4,1fr);border:1px solid #E3E6E3;margin:8px 0 14px}.kf div{padding:10px 12px;border-right:1px solid #E3E6E3}.kf div:last-child{border-right:0}.kf b{display:block;font-size:17pt;color:#1E7A44}.kf span{font-size:8.5pt;color:#4B5550}'''

def q_html():
    out = ['<h1>PlateLoop questionnaires</h1><p class="sub">Four short surveys: students, canteen and kitchen staff, hospital staff, and office workers. Each takes about four minutes and is anonymous.</p>']
    for G in GROUPS:
        out.append(f'<section class="grp"><h2>{G["title"]}</h2><p class="sub">{G["intro"]}</p>')
        for code, text, typ, opts, _ in G['qs']:
            if typ == 'cat': body = '<div class="opts">' + ''.join(f'<span><i></i>{o}</span>' for o in opts) + '</div>'
            elif typ == 'lik': body = f'<div class="scale">{opts[0]} ' + ''.join(f'<i>{v}</i>' for v in range(1, 6)) + f' {opts[1]}</div>'
            elif typ == 'num': body = '<div class="line"></div>'
            else: body = '<div class="box"></div>'
            out.append(f'<div class="q"><b><span>{code}</span>{text}</b>{body}</div>')
        out.append('</section>')
    return f'<!doctype html><html><head><meta charset="utf-8"><title>PlateLoop questionnaires</title><link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Inter:wght@400;600;800&display=swap"><style>{CSS}</style></head><body>{"".join(out)}</body></html>'

def bars(items, fmt=lambda v: f'{P(v)}', width=330, color='#1E7A44', mx=None):
    mx = mx or max(v for _, v in items) or 1; rh = 22; H = len(items) * rh + 4; lw = 150
    s = ''.join(f'<text x="0" y="{i*rh+15}" font-size="9.5" fill="#4B5550">{lab}</text><rect x="{lw}" y="{i*rh+4}" width="{(width-lw-50)*v/mx:.1f}" height="14" fill="{color}"/><text x="{lw+(width-lw-50)*v/mx+5:.1f}" y="{i*rh+15}" font-size="9.5" font-weight="700" fill="#16201A">{fmt(v)}</text>' for i, (lab, v) in enumerate(items))
    return f'<svg viewBox="0 0 {width} {H}" width="100%" font-family="Inter, Arial">{s}</svg>'

def findings_html(data, texts):
    S, K, Hh, O = (data[k] for k in ('Students', 'Kitchen', 'Hospital', 'Office'))
    share = lambda df, c, o: (df[c] == o).mean()
    agree = lambda df, c: (df[c] >= 4).mean()
    charts = [
        ('Students: main reason for leaving food', bars([(o, share(S, 'S3', o)) for o in GROUPS[0]['qs'][2][3]])),
        ('Students: what they would rather see', bars([(o, share(S, 'S7', o)) for o in GROUPS[0]['qs'][6][3]])),
        ('Kitchen staff: how quantities are decided', bars([(o, share(K, 'K2', o)) for o in GROUPS[1]['qs'][1][3]])),
        ('Hospital staff: when a patient who stops eating is noticed', bars([(o, share(Hh, 'H5', o)) for o in GROUPS[2]['qs'][4][3]], color='#2E64B0')),
        ('Office workers: share of meal data with employer', bars([(o, share(O, 'O8', o)) for o in GROUPS[3]['qs'][7][3]], color='#5046C8')),
        ('Share who agree (4–5) by statement', bars([('Kitchen: suggested order', agree(K, 'K6')), ('Kitchen: waste by dish', agree(K, 'K7')), ('Hospital: automatic intake', agree(Hh, 'H7')), ('Hospital: alert with reason', agree(Hh, 'H8')), ('Office: daily pick for goal', agree(O, 'O7')), ('Students: class competition', agree(S, 'S9'))], width=330)),
    ]
    kf = f'''<div class="kf"><div><b>{P(share(S,'S3','Portion too big'))}</b><span>of students leave food mainly because the portion is too big</span></div>
      <div><b>{P(share(K,'K4','Yes, accurately'))}</b><span>of kitchen staff know accurately which dishes come back</span></div>
      <div><b>{P(agree(Hh,'H3'))}</b><span>of hospital staff trust their intake records</span></div>
      <div><b>{P(share(O,'O8','Only anonymous totals')+share(O,'O8','No'))}</b><span>of office workers want meal data kept private or anonymous</span></div></div>'''
    figs = ''.join(f'<div><b style="font-size:10pt">{t}</b>{svg}</div>' for t, svg in charts)
    fds = ''.join(f'<div class="fd"><span class="n">{i}</span><div>{t}<em>→ {m}</em></div></div>' for i, (t, m) in enumerate(texts, 1))
    n = ', '.join(f"{G['key'].lower()} {G['n']}" for G in GROUPS)
    return f'''<!doctype html><html><head><meta charset="utf-8"><title>PlateLoop survey key findings (demo)</title><link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Inter:wght@400;600;800&display=swap"><style>{CSS}</style></head><body>
    <h1>PlateLoop survey: key findings</h1><p class="sub">Students, canteen and kitchen staff, hospital staff and office workers.</p>
    <div class="demo"><b>Demo data.</b> These results come from synthetic responses ({n}) generated to demonstrate the analysis. They are not real survey results and should not be quoted as evidence. The questionnaires and analysis workbook are ready for real responses.</div>
    {kf}<div class="grid2">{figs}</div><p class="cap">Figures computed from the demo responses in plateloop-survey-analysis-DEMO.xlsx.</p>
    <h2 style="margin-top:12px">Findings and what they mean for PlateLoop</h2>{fds}</body></html>'''

def finding_texts(data):
    """The same findings as the workbook's formulas, computed in Python for the PDF."""
    S, K, Hh, O = (data[k] for k in ('Students', 'Kitchen', 'Hospital', 'Office'))
    sh = lambda df, c, o: f'{P((df[c] == o).mean())}'; ag = lambda df, c: f'{P((df[c] >= 4).mean())}'
    m = lambda s: f'{s.mean():.1f}'
    sec, uni = S[S.S1 == 'Secondary school'], S[S.S1 == 'University']
    return [
     (f'Portion size is the top reason for leaving food: {sh(S,"S3","Portion too big")} chose it, and {sh(S,"S4","Yes")} would take a smaller portion if they could ask easily.', 'The before scan and right-portion rewards in Loopi: a small plate finished counts fully.'),
     (f'More students prefer food groups ({sh(S,"S7","Food groups")}) than calories ({sh(S,"S7","Calories")}).', 'Loopi shows food groups by default; calories stay optional.'),
     (f'A virtual pet appeals more to secondary students (average {m(sec.S8)}/5) than to university students ({m(uni.S8)}/5).', 'Lead with the Loopi pet in schools; emphasise nutrition feedback on campuses.'),
     (f'Only {ag(S,"S10")} of students are comfortable with face sign-in.', 'Keep the class-and-number alternative; face sign-in stays optional.'),
     (f'Most kitchen staff plan by experience or headcount ({P(((K.K2=="Past experience")|(K.K2=="Headcount × standard portion")).mean())}), and only {sh(K,"K4","Yes, accurately")} know accurately which dishes come back.', 'Kitchen dashboard: waste by dish and a forecast from real eating data.'),
     (f'Staff estimate {K.K3.mean():.0f}% of cooked food is thrown away, and spend {K.K9.mean():.0f} minutes a day planning quantities.', 'A suggested order for tomorrow saves planning time and cuts over-preparation.'),
     (f'A suggested order is rated useful by {ag(K,"K6")} and waste by dish by {ag(K,"K7")}, but {ag(K,"K10")} worry about extra workload.', 'Measurement is automatic: no forms or manual entry for kitchen staff.'),
     (f'Only {ag(Hh,"H3")} of hospital staff trust their intake records, and they spend {Hh.H4.mean():.0f} minutes a shift documenting intake.', 'Loopi Care records every tray automatically.'),
     (f'Patients who stop eating are noticed the same day in only {sh(Hh,"H5","Same day")} of cases.', 'Nurse-station alerts after poor meals, the same day.'),
     (f'Automatic intake records are rated useful by {ag(Hh,"H7")}; {sh(Hh,"H9","Yes")} say patients would benefit from choosing tomorrow’s meals.', 'Loopi Care: per-tray intake and patient meal choice.'),
     (f'Only {ag(O,"O4")} of office workers know the nutrients in their lunch, and {P(((O.O5=="Often")|(O.O5=="Daily")).mean())} feel an afternoon energy dip often or daily.', 'Loopi Work: nutrients per lunch and a 3 pm energy check-in.'),
     (f'{sh(O,"O6","Yes")} would pre-order lunch to skip a typical {O.O9.mean():.0f}-minute queue.', 'Pre-ordering in Loopi Work; kitchens cook to orders.'),
     (f'Only {sh(O,"O8","Yes")} would share meal data with their employer; {sh(O,"O8","Only anonymous totals")} would share anonymous totals only.', 'Private by default: employers see team totals, never names.'),
    ]

def to_pdf(html_path, pdf_path):
    sys.path.insert(0, os.path.join(ROOT, 'tools'))
    from capture import serve, PORT
    from selenium import webdriver
    h = serve(); o = webdriver.ChromeOptions(); o.add_argument('--headless=new'); d = webdriver.Chrome(options=o)
    try:
        rel = os.path.relpath(html_path, ROOT).replace(os.sep, '/')
        d.get(f'http://127.0.0.1:{PORT}/{rel}'); time.sleep(2.5)
        pdf = d.execute_cdp_cmd('Page.printToPDF', {'printBackground': True, 'preferCSSPageSize': True})
        open(pdf_path, 'wb').write(base64.b64decode(pdf['data']))
    finally:
        d.quit(); h.shutdown()

def main():
    os.makedirs(OUT, exist_ok=True)
    data = make_data()
    build_workbook(data, os.path.join(OUT, 'plateloop-survey-analysis-DEMO.xlsx'))
    texts = finding_texts(data)
    for name, html in (('questionnaires', q_html()), ('key-findings-DEMO', findings_html(data, texts))):
        hp = os.path.join(OUT, f'plateloop-{name}.html'); open(hp, 'w', encoding='utf-8').write(html)
        to_pdf(hp, os.path.join(OUT, f'plateloop-{name}.pdf'))
    print('done:', os.listdir(OUT))
    for t, _ in texts: print(' -', t)

if __name__ == '__main__':
    main()
